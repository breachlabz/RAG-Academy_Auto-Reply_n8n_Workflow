"""Turn a .docx into the heading-marked markdown the chunker expects.

Written against a document that carries no heading styles at all -- every
paragraph is `Normal`, headings are marked only by being entirely bold, and the
prose is hard-wrapped so single sentences arrive split across three or four
paragraphs. Both are typical of a Word file that started life as a designed
layout rather than a structured document, so the conversion has to recover
structure that Word never recorded.

Two things this deliberately does *not* do:

- It ignores images. A 2.8MB file holding 8KB of text is mostly pictures, and
  whatever those pictures say is invisible to retrieval. If a diagram carries
  content you expect answers about, that content has to be typed out somewhere.
- It reads only paragraphs and tables. Text inside shapes and text boxes
  (`w:txbxContent`) is not picked up; `verify()` reports the shortfall so that
  failure is visible rather than silent.
"""

from __future__ import annotations

import pathlib
import re
import zipfile

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

# A fully bold paragraph is a heading; a partly bold one is body text with a
# bold lead-in ("**Level 1** gives the participant ..."), which is a different
# thing entirely and must not become a heading.
MAX_HEADING_CHARS = 90


def _iter_blocks(document: DocxDocument):
    """Paragraphs and tables in document order.

    python-docx exposes `.paragraphs` and `.tables` as separate sequences,
    which loses the interleaving -- and the interleaving is the only thing that
    says which heading a table belongs under.
    """
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _is_heading(para: Paragraph) -> bool:
    runs = [r for r in para.runs if r.text.strip()]
    if not runs or len(para.text.strip()) > MAX_HEADING_CHARS:
        return False
    return all(r.bold for r in runs)


def _table_markdown(table: Table) -> list[str]:
    """Render a table as one self-describing line per row.

    A pipe table embeds badly: the column headers end up in a different chunk
    from most of the rows, so a row reading "Cryptography | Differentiation of
    cryptographic functions" loses the fact that the first column is a module
    name. Repeating the header per row costs tokens and keeps every row
    meaningful on its own.
    """
    rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in table.rows]
    rows = [r for r in rows if any(cell for cell in r)]
    if not rows:
        return []

    header, *body = rows
    if not body:  # Single-row table: nothing to label, emit it plainly.
        return [" — ".join(header)]

    out = []
    for row in body:
        pairs = [
            f"{head}: {cell}"
            for head, cell in zip(header, row)
            if cell
        ]
        out.append(re.sub(r"\s+", " ", " — ".join(pairs)).strip())
    return out


def to_markdown(path: pathlib.Path) -> str:
    """Convert one .docx to markdown with `#`/`##` headings."""
    document = docx.Document(str(path))
    lines: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        """Re-join hard-wrapped paragraphs into one flowing block."""
        if buffer:
            lines.append(re.sub(r"\s+", " ", " ".join(buffer)).strip())
            lines.append("")
            buffer.clear()

    for block in _iter_blocks(document):
        if isinstance(block, Table):
            flush()
            for row in _table_markdown(block):
                lines.append(row)
                lines.append("")
            continue

        text = block.text.strip()
        if not text:
            flush()
            continue

        style = block.style.name
        if style == "Title":
            flush()
            lines += [f"# {text}", ""]
        elif style == "Subtitle" or _is_heading(block):
            flush()
            lines += [f"## {text}", ""]
        else:
            buffer.append(text)

    flush()
    # A heading immediately followed by another heading usually means a label
    # for the block below it; harmless, but collapse the blank runs.
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"


def verify(path: pathlib.Path) -> dict[str, int]:
    """Compare extracted characters against every text run in the file.

    Silent under-extraction is the failure mode that matters: text living in a
    shape or a text box simply never appears, and the only symptom is the model
    refusing questions it should be able to answer.
    """
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf8")
        images = sum(1 for n in archive.namelist() if n.startswith("word/media"))

    # Measured on the raw blocks, not on to_markdown's output: the converter
    # repeats a table's header on every row, so its output legitimately runs
    # longer than the source and would mask a shortfall rather than show it.
    document = docx.Document(str(path))
    reached = 0
    for block in _iter_blocks(document):
        if isinstance(block, Table):
            reached += sum(len(c.text) for r in block.rows for c in r.cells)
        else:
            reached += len(block.text)

    in_file = sum(len(t) for t in re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))
    return {
        "chars_in_file": in_file,
        "chars_reached": reached,
        "textboxes": xml.count("txbxContent"),
        "images": images,
    }
