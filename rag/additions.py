"""New knowledge added from the UI, kept as real documents: data/docs/Additions_N.docx.

The knowledge tab no longer edits or deletes existing chunks -- the source
.docx files are the only way content changes. Adding is the one write left, and
it goes through the same path as every other document: the entry is appended to
an `Additions_N.docx` next to the rest of the corpus, and that file is ingested
like any other. So an addition gets the same id scheme (`Additions_1.docx#3`),
the same metadata (`source`, `heading`, `chunk_index`) and the same row in
`knowledge_chunks`, and it survives `python -m rag ingest --reset` because it is
just another file in data/docs.

Each file holds at most MAX_PER_DOC entries; the next one starts
`Additions_<N+1>.docx`. An entry is a Subtitle-styled heading (which
rag.docx_text reads as `##`, i.e. a chunk boundary) followed by its paragraphs,
so each addition is its own chunk. Entries are only ever appended, which keeps
the chunk ids of earlier entries stable.

If the embed or Chroma write fails, the file is put back the way it was, so the
document and the vector store never disagree.
"""

from __future__ import annotations

import os
import pathlib
import re
import threading

MAX_PER_DOC = 20
PREFIX = "Additions_"
HEADING_STYLE = "Subtitle"
MAX_TITLE_CHARS = 90

# One writer at a time: two concurrent adds would both pick the same file and
# the second save would drop the first entry.
_lock = threading.Lock()


def _docs_dir() -> pathlib.Path:
    from .core import DOCS_DIR

    return DOCS_DIR


def _files(docs_dir: pathlib.Path) -> list[tuple[int, pathlib.Path]]:
    found = []
    for path in docs_dir.glob(f"{PREFIX}*.docx"):
        match = re.fullmatch(rf"{PREFIX}(\d+)\.docx", path.name)
        if match:
            found.append((int(match.group(1)), path))
    return sorted(found)


def entry_count(path: pathlib.Path) -> int:
    import docx

    return sum(1 for p in docx.Document(str(path)).paragraphs if p.style.name == HEADING_STYLE)


def _target(docs_dir: pathlib.Path) -> tuple[pathlib.Path, int]:
    """The file the next entry goes into, and how many entries it already has."""
    files = _files(docs_dir)
    if files:
        number, path = files[-1]
        count = entry_count(path)
        if count < MAX_PER_DOC:
            return path, count
        return docs_dir / f"{PREFIX}{number + 1}.docx", 0
    return docs_dir / f"{PREFIX}1.docx", 0


def _title(title: str, content: str) -> str:
    title = " ".join((title or "").split())
    if not title:
        # No title given: the first line of the content stands in, so the
        # chunk still carries a meaningful heading trail into its embedding.
        title = " ".join(content.strip().splitlines()[0].split())
    if len(title) > MAX_TITLE_CHARS:
        title = title[: MAX_TITLE_CHARS - 1].rstrip() + "…"
    return title


def add(content: str, title: str = "", *, docs_dir: pathlib.Path | None = None) -> dict:
    """Append one entry to the current Additions doc and ingest that file.

    Returns {"file", "entry", "max", "chunks"}: the file name, this entry's
    1-based position in it, MAX_PER_DOC, and the knowledge rows the entry
    produced. ValueError for empty content; embed/Chroma errors propagate after
    the file has been restored.
    """
    import docx

    from . import knowledge
    from .core import ingest

    if not isinstance(content, str) or not content.strip():
        raise ValueError("content cannot be empty")
    docs_dir = pathlib.Path(docs_dir) if docs_dir else _docs_dir()
    heading = _title(title, content)

    with _lock:
        docs_dir.mkdir(parents=True, exist_ok=True)
        path, count = _target(docs_dir)
        previous = path.read_bytes() if path.exists() else None
        before = {c["id"] for c in knowledge.list_chunks(source=path.name)}

        document = docx.Document(str(path)) if previous is not None else docx.Document()
        document.add_paragraph(heading, style=HEADING_STYLE)
        # One paragraph per line with a blank paragraph after each: docx_text
        # joins consecutive paragraphs into one block, which would collapse a
        # list into a single run-on line.
        for line in content.strip().splitlines():
            if line.strip():
                document.add_paragraph(line.strip())
                document.add_paragraph("")
        document.save(str(path))
        if previous is None:
            # The API runs as root in its container; hand the new file to
            # whoever owns data/docs so it stays editable on the host.
            try:
                st = docs_dir.stat()
                os.chown(path, st.st_uid, st.st_gid)
                os.chmod(path, 0o664)
            except OSError:
                pass

        try:
            ingest([path])
        except Exception:
            if previous is None:
                path.unlink(missing_ok=True)
            else:
                path.write_bytes(previous)
            raise

    chunks = [c for c in knowledge.list_chunks(source=path.name) if c["id"] not in before]
    return {"file": path.name, "entry": count + 1, "max": MAX_PER_DOC, "chunks": chunks}
